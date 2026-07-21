"""Excel and KeyLab SQL helpers used by the SQL-only progress pipeline."""

import os
import re
import subprocess
import sys
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import load_workbook


BASE_DIR = Path(__file__).parent.resolve()
DATA_DIR = BASE_DIR / "Data"
CLEAN_DIR = BASE_DIR / "File_sach"

SHEET_HINTS = ["Đơn hàng", "Don hang", "Sheet1", "Sheet"]
COL_HINTS = ["Mã ĐH", "mã_dh", "ma_dh", "Mã đơn", "MaDH", "ORDER_ID"]

_NO_WINDOW = subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0


def log(message):
    print(message, flush=True)


def normalize_text(value: Optional[str]) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def normalize_ma_dh(value) -> str:
    order_id = normalize_text(value)
    if re.fullmatch(r"\d+\.0", order_id):
        order_id = order_id[:-2]
    return order_id


def detect_sheet_col(xlsx_path: str):
    workbook = pd.ExcelFile(xlsx_path)
    sheets = workbook.sheet_names
    sheet = next((name for name in SHEET_HINTS if name in sheets), sheets[0])
    columns = list(workbook.parse(sheet_name=sheet, nrows=1).columns)

    def normalize(value):
        return str(value).lower().replace(" ", "").replace("_", "").replace("đ", "d")

    normalized_hints = [normalize(hint) for hint in COL_HINTS]
    column = next((item for item in columns if normalize(item) in normalized_hints), columns[0])
    log(f'[keylab] Sheet: "{sheet}" | Cột mã đơn: "{column}"')
    return sheet, column


def load_order_ids(xlsx_path: str, sheet_name: str, column_name: str) -> list[str]:
    workbook = pd.ExcelFile(xlsx_path)
    if sheet_name not in workbook.sheet_names:
        raise ValueError(f"Không tìm thấy sheet '{sheet_name}' trong file Excel")

    frame = workbook.parse(sheet_name=sheet_name)
    if column_name not in frame.columns:
        raise KeyError(f"Không tìm thấy cột '{column_name}' trong sheet '{sheet_name}'")

    seen = set()
    order_ids = []
    for value in frame[column_name].tolist():
        order_id = normalize_ma_dh(value)
        if order_id and order_id not in seen:
            order_ids.append(order_id)
            seen.add(order_id)
    return order_ids


def build_order_summary(progress_df: pd.DataFrame) -> pd.DataFrame:
    """Build compatibility columns consumed by the existing Excel cleaner."""
    if progress_df.empty:
        return pd.DataFrame(columns=[
            "ma_dh", "so_cd_cao_web", "ktv_theo_cong_doan_live",
            "ds_cong_doan_live", "lan_cap_nhat_cao_web", "tai_khoan_cao",
        ])

    def aggregate_ktv(group: pd.DataFrame) -> str:
        values = []
        for _, row in group.iterrows():
            stage = normalize_text(row.get("cong_doan"))
            technician = normalize_text(row.get("ten_ktv"))
            confirmed = normalize_text(row.get("xac_nhan"))
            if not stage and not technician:
                continue
            state = "Có" if technician not in ("", "-") and confirmed == "Có" else "Chưa"
            value = f"{stage}: {state}".strip(": ")
            if value and value not in values:
                values.append(value)
        return " | ".join(values)

    def aggregate_unique(group: pd.DataFrame, column: str, separator: str) -> str:
        values = []
        for raw_value in group[column].tolist():
            value = normalize_text(raw_value)
            if value and value not in values:
                values.append(value)
        return separator.join(values)

    return progress_df.groupby("ma_dh", as_index=False).apply(
        lambda group: pd.Series({
            "so_cd_cao_web": len(group),
            "ktv_theo_cong_doan_live": aggregate_ktv(group),
            "ds_cong_doan_live": aggregate_unique(group, "cong_doan", " -> "),
            "lan_cap_nhat_cao_web": pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S"),
            "tai_khoan_cao": aggregate_unique(group, "tai_khoan_cao", " | "),
        })
    ).reset_index(drop=True)


def _autosize_sheet(worksheet, max_width: int = 45) -> None:
    from openpyxl.utils import get_column_letter

    for column_cells in worksheet.iter_cols():
        cells = list(column_cells)
        if not cells:
            continue
        letter = get_column_letter(cells[0].column)
        width = 10
        for cell in cells[:200]:
            value = "" if cell.value is None else str(cell.value)
            width = max(width, min(max_width, len(value) + 2))
        worksheet.column_dimensions[letter].width = width


def merge_back_to_workbook(input_xlsx: str, output_xlsx: str, progress_df: pd.DataFrame) -> None:
    """Add SQL progress using legacy sheet names required by the current importer."""
    workbook = load_workbook(input_xlsx)

    if "Cao_web_live" in workbook.sheetnames:
        del workbook["Cao_web_live"]
    raw_sheet = workbook.create_sheet("Cao_web_live")
    raw_headers = list(progress_df.columns) if not progress_df.empty else [
        "ma_dh", "thu_tu", "cong_doan", "ten_ktv", "xac_nhan",
        "thoi_gian_hoan_thanh", "raw_row_text", "tai_khoan_cao", "barcode_labo",
    ]
    raw_sheet.append(raw_headers)
    if not progress_df.empty:
        for row in progress_df[raw_headers].itertuples(index=False, name=None):
            raw_sheet.append(list(row))
    _autosize_sheet(raw_sheet)

    summary_df = build_order_summary(progress_df)
    extra_columns = [
        "ktv_theo_cong_doan_live", "ds_cong_doan_live", "so_cd_cao_web",
        "lan_cap_nhat_cao_web", "tai_khoan_cao",
    ]
    for sheet_name in ("Tong_hop_don_hang", "Kanban_ready"):
        if sheet_name not in workbook.sheetnames:
            continue
        worksheet = workbook[sheet_name]
        headers = {
            str(worksheet.cell(1, column).value): column
            for column in range(1, worksheet.max_column + 1)
            if worksheet.cell(1, column).value
        }
        if "ma_dh" not in headers:
            continue
        for column_name in extra_columns:
            if column_name not in headers:
                headers[column_name] = worksheet.max_column + 1
                worksheet.cell(1, headers[column_name]).value = column_name
        lookup = {
            normalize_ma_dh(row["ma_dh"]): row.to_dict()
            for _, row in summary_df.iterrows()
        }
        for row_index in range(2, worksheet.max_row + 1):
            order_id = normalize_ma_dh(worksheet.cell(row_index, headers["ma_dh"]).value)
            item = lookup.get(order_id)
            if not item:
                continue
            for column_name in extra_columns:
                worksheet.cell(row_index, headers[column_name]).value = item.get(column_name, "")
        _autosize_sheet(worksheet)

    if "Cong_doan_chi_tiet_live" in workbook.sheetnames:
        del workbook["Cong_doan_chi_tiet_live"]
    detail_sheet = workbook.create_sheet("Cong_doan_chi_tiet_live")
    detail_headers = [
        "ma_dh", "thu_tu", "cong_doan", "ten_ktv", "xac_nhan",
        "thoi_gian_hoan_thanh", "raw_row_text", "tai_khoan_cao",
    ]
    detail_sheet.append(detail_headers)
    if not progress_df.empty:
        for row in progress_df[detail_headers].itertuples(index=False, name=None):
            detail_sheet.append(list(row))
    _autosize_sheet(detail_sheet)
    workbook.save(output_xlsx)


def sync_keylab_sql_notes(order_ids, stem: str) -> bool:
    if os.environ.get("KEYLAB_SQL_NOTES_DISABLED", "").strip() == "1":
        log("[keylab] SQL notes disabled by environment")
        return False

    script = BASE_DIR / "keylab_sql_notes_scraper.ps1"
    if not script.exists():
        log("[keylab] SQL notes script not found; skipped")
        return False
    if sys.platform != "win32":
        log("[keylab] SQL notes require Windows PowerShell; skipped")
        return False

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    order_file = DATA_DIR / f"{stem}_keylab_order_ids.txt"
    try:
        order_file.write_text(
            "\n".join(str(value).strip() for value in order_ids if str(value).strip()),
            encoding="utf-8",
        )
        result = subprocess.run(
            [
                "powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass",
                "-File", str(script), "-OrderFile", str(order_file),
                "-OutFile", str(BASE_DIR / "keylab_notes.json"),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=180,
            creationflags=_NO_WINDOW,
        )
        output = ((result.stdout or "") + (result.stderr or "")).strip()
        if result.returncode == 0:
            log(f"[keylab] SQL notes OK: {output}")
            return True
        log(f"[keylab] SQL notes failed (code {result.returncode}): {output[-600:]}")
    except Exception as exc:
        log(f"[keylab] SQL notes error: {exc}")
    finally:
        order_file.unlink(missing_ok=True)
    return False
